"""S3.3 tests: archive eligibility, ledger export, manual review flow."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from config.schema import InvoiceConfig
from invoice.eligibility import evaluate_eligibility
from invoice.ledger_exporter import MonthlyLedgerExporter
from models.invoice import ArchiveCandidate, InvoiceRecord
from services.invoice_service import InvoiceService, plan_month
from utils.paths import RuntimePaths


def _record(**kw) -> InvoiceRecord:
    record = InvoiceRecord(
        file_name="inv.pdf",
        file_path="/tmp/inv.pdf",
        invoice_number="26512000000000000001",
        invoice_date="2026-10-08",
        amount=99.0,
        tax_amount=11.39,
        buyer="甲城购买方科技有限公司",
        seller="乙郡销售方服务有限公司",
        invoice_content="*餐饮服务*餐费",
        doc_type="e_invoice",
        file_hash="abc123",
        recognition_status="complete",
    )
    record.field_sources = {
        "number": "dzfp_block",
        "date": "dzfp_block",
        "amount": "verified",
        "tax": "heji_pair",
        "buyer": "order_pair",
        "seller": "order_pair",
    }
    for key, value in kw.items():
        setattr(record, key, value)
    return record


# ----------------------------------------------------------------------
# Eligibility


def test_eligible_clean_record() -> None:
    verdict = evaluate_eligibility(_record(), "新发票")
    assert verdict.eligible
    assert verdict.reasons == []


def test_complete_alone_is_not_enough() -> None:
    """A 'complete' record with only fallback evidence must be blocked."""
    record = _record()
    record.field_sources = {"number": "fallback", "date": "fallback", "amount": "fallback"}
    verdict = evaluate_eligibility(record, "新发票")
    assert not verdict.eligible
    assert {"weak_number", "weak_date", "weak_amount", "tax_missing"} <= set(verdict.reasons)


def test_non_invoice_blocked_even_when_complete() -> None:
    record = _record(doc_type="hotel_folio")
    assert not evaluate_eligibility(record, "新发票").eligible


def test_duplicate_blocked() -> None:
    for status in ("强重复", "疑似重复", "待确认"):
        assert not evaluate_eligibility(_record(), status).eligible


def test_review_flags_block() -> None:
    record = _record()
    record.review_fields = ["buyer"]
    verdict = evaluate_eligibility(record, "新发票")
    assert not verdict.eligible
    assert "review_flags" in verdict.reasons


def test_missing_seller_blocks() -> None:
    record = _record(seller="")
    verdict = evaluate_eligibility(record, "新发票")
    assert not verdict.eligible
    assert "seller_missing" in verdict.reasons


def test_bad_number_format_blocks() -> None:
    record = _record(invoice_number="12-34")
    verdict = evaluate_eligibility(record, "新发票")
    assert not verdict.eligible
    assert "bad_number_format" in verdict.reasons


def test_hash_mismatch_blocks() -> None:
    assert not evaluate_eligibility(_record(), "新发票", file_hash_ok=False).eligible


# ----------------------------------------------------------------------
# Service dry-run on fixture PDFs


def _service(tmp_path: Path) -> tuple[InvoiceService, InvoiceConfig, Path]:
    paths = RuntimePaths.from_root(tmp_path / "state")
    config = InvoiceConfig()
    root = tmp_path / "发票管理"
    pending = root / config.pending_dir_name
    pending.mkdir(parents=True)
    config.root_dir = str(root)
    return InvoiceService(paths), config, pending


def _write_invoice_pdf(path: Path, text: str) -> None:
    from invoice.pdf_text_extractor import PdfTextExtractor

    del text  # placeholder — real fixtures below write binary PDFs


def _make_pdf(path: Path, body: str) -> None:
    """Minimal single-page PDF with the given text stream."""
    content = f"BT /F1 10 Tf 10 700 Td ({body}) Tj ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = ["%PDF-1.4"]
    for idx, obj in enumerate(objects, start=1):
        out.append(f"{idx} 0 obj {obj} endobj")
    out.append("trailer << /Root 1 0 R >>")
    path.write_bytes("\n".join(out).encode("latin-1"))


def test_evaluate_pending_is_read_only(tmp_path: Path, monkeypatch) -> None:
    service, config, pending = _service(tmp_path)
    _make_pdf(pending / "a.pdf", "hello")
    before = sorted(p.name for p in pending.iterdir())
    candidates = service.evaluate_pending(config)
    assert len(candidates) == 1
    assert not candidates[0].eligible
    assert sorted(p.name for p in pending.iterdir()) == before


def test_auto_archive_dry_run_moves_nothing(tmp_path: Path) -> None:
    service, config, pending = _service(tmp_path)
    _make_pdf(pending / "a.pdf", "hello")
    candidates = service.auto_archive(config, dry_run=True)
    assert len(candidates) == 1
    assert (pending / "a.pdf").exists()


# ----------------------------------------------------------------------
# Manual overrides survive rescan


def test_manual_overrides_survive_rescan(tmp_path: Path) -> None:
    service, config, pending = _service(tmp_path)
    _make_pdf(pending / "a.pdf", "hello")
    candidate = service.evaluate_pending(config)[0]
    assert not candidate.eligible  # garbage pdf -> not complete anyway
    record = candidate.record
    service.save_manual_review(record, {"invoice_number": "26512000000000000099"})

    again = service.evaluate_pending(config)[0]
    assert again.record.invoice_number == "26512000000000000099"


# ----------------------------------------------------------------------
# Monthly ledger exporter


def _ledger_record(**kw) -> InvoiceRecord:
    record = _record(archive_status="auto_archived", review_status="confirmed")
    record.archived_at = "2026-10-20T10:00:00"
    for key, value in kw.items():
        setattr(record, key, value)
    return record


def test_ledger_two_sheets_and_totals(tmp_path: Path) -> None:
    exporter = MonthlyLedgerExporter()
    rows = [
        _ledger_record(amount=99.0, invoice_date="2026-10-08"),
        _ledger_record(amount=720.0, invoice_date="2026-10-02",
                       invoice_number="26512000000000000002"),
        _ledger_record(amount=306.0, invoice_date="2026-10-19",
                       invoice_number="26512000000000000003"),
    ]
    out = exporter.export(tmp_path / "月度台账_202610.xlsx", rows, reimburser_default="张三")

    book = openpyxl.load_workbook(out)
    flow = book["报销流水"]
    assert [c.value for c in flow[1]] == [
        "序号", "日期", "姓名", "金额", "类别", "凭证类型", "费用发生地", "备注",
    ]
    # Sorted by date ascending
    assert [flow.cell(row=r, column=2).value for r in (2, 3, 4)] == [
        "2026.10.02", "2026.10.08", "2026.10.19",
    ]
    assert flow.cell(row=2, column=3).value == "张三"
    assert flow.cell(row=2, column=4).value == 720.0
    assert flow.cell(row=2, column=8).value in (None, "")
    total_row = flow.max_row
    assert flow.cell(row=total_row, column=4).value == 1125.0
    assert flow.auto_filter.ref
    assert flow.freeze_panes == "A2"

    detail = book["发票明细"]
    headers = [c.value for c in detail[1]]
    assert "报销状态" in headers and "发票号码" in headers
    assert detail.cell(row=2, column=2).value == "26512000000000000002"
    assert detail.cell(row=2, column=11).value == "自动归档"
    assert detail.cell(row=2, column=headers.index("报销状态") + 1).value == "未确定"


def test_ledger_regeneration_does_not_duplicate(tmp_path: Path) -> None:
    exporter = MonthlyLedgerExporter()
    rows = [_ledger_record(amount=99.0)]
    out = tmp_path / "月度台账_202610.xlsx"
    exporter.export(out, rows)
    exporter.export(out, rows)  # regenerate
    book = openpyxl.load_workbook(out)
    flow = book["报销流水"]
    data_rows = [
        r for r in range(2, flow.max_row + 1)
        if isinstance(flow.cell(row=r, column=1).value, int)
    ]
    assert len(data_rows) == 1


def test_export_monthly_ledger_filters_archived(tmp_path: Path) -> None:
    service, config, _pending = _service(tmp_path)
    store = service.metadata_store
    archived = _ledger_record()
    pending_rec = _record(archive_status="")
    store.save_records([archived, pending_rec])
    out = service.export_monthly_ledger("202610", config)
    book = openpyxl.load_workbook(out)
    assert book["报销流水"].max_row == 3  # header + 1 row + total


def test_plan_month() -> None:
    assert plan_month("2026-10-08") == "202610"
    assert plan_month("") == ""
    assert plan_month("2026-10") == ""


def test_review_summary_splits_buckets(tmp_path: Path) -> None:
    service, _config, _pending = _service(tmp_path)

    def cand(doc_type: str, eligible: bool = False) -> ArchiveCandidate:
        return ArchiveCandidate(
            record=_record(doc_type=doc_type),
            source_path="/tmp/x.pdf",
            duplicate_status="新发票",
            duplicate_reason="",
            eligible=eligible,
            reasons=[] if eligible else ["not_e_invoice"],
        )

    candidates = [
        cand("e_invoice", eligible=True),
        cand("e_invoice"),  # e_invoice with problems -> needs human review
        cand("hotel_folio"),
        cand("travel_voucher"),
        cand("needs_ocr"),
        cand("unknown"),
    ]
    summary = service.review_summary(candidates)
    assert summary == {
        "eligible": 1,
        "needs_review": 1,
        "other_voucher": 2,
        "unidentified": 2,
    }


def test_reviewed_record_is_not_self_duplicate(tmp_path: Path) -> None:
    """Regression: save_manual_review stores the record; the next dedup
    must not flag it as a strong duplicate of itself (S3.3 first-archive
    bug: 确认归档 failed with '强重复(matching file hash)')."""
    service, config, pending = _service(tmp_path)
    _make_pdf(pending / "a.pdf", "hello")
    record = service.evaluate_pending(config)[0].record
    service.save_manual_review(
        record, {"invoice_number": "26512000000000000099", "amount": 88.0}
    )
    again = service.evaluate_pending(config)[0]
    assert again.duplicate_status != "强重复"
    # A DIFFERENT file with the same hash is still a real duplicate.
    import shutil
    shutil.copy2(pending / "a.pdf", pending / "copy.pdf")
    dup = service.evaluate_pending(config)
    by_name = {c.record.file_name: c for c in dup}
    assert by_name["copy.pdf"].duplicate_status == "强重复"


def test_archive_non_invoices_moves_and_keeps(tmp_path: Path) -> None:
    service, config, pending = _service(tmp_path)

    def cand(name: str, doc_type: str) -> ArchiveCandidate:
        return ArchiveCandidate(
            record=_record(doc_type=doc_type, file_name=name,
                           file_path=str(pending / name)),
            source_path=str(pending / name),
            duplicate_status="新发票", duplicate_reason="",
            eligible=False, reasons=["not_e_invoice"],
        )

    (pending / "folio.pdf").write_bytes(b"A" * 10)
    (pending / "rental.pdf").write_bytes(b"B" * 10)
    (pending / "scan.pdf").write_bytes(b"C" * 10)
    (pending / "misc.pdf").write_bytes(b"D" * 10)

    result = service.archive_non_invoices(
        [cand("folio.pdf", "hotel_folio"),
         cand("rental.pdf", "unknown"),
         cand("scan.pdf", "needs_ocr"),
         cand("misc.pdf", "other_pdf")],
        config,
    )
    review = Path(config.root_dir) / "非发票附件"
    assert (review / "folio.pdf").exists()
    assert (review / "misc.pdf").exists()
    # 不确定/低置信度的留下等人工：unknown 和 needs_ocr 都保留
    assert (pending / "scan.pdf").exists()
    assert (pending / "rental.pdf").exists()
    assert not (pending / "folio.pdf").exists()
    assert len(result["moved"]) == 2
    assert len(result["kept"]) == 2


def test_archive_non_invoices_merges_same_hash(tmp_path: Path) -> None:
    service, config, pending = _service(tmp_path)
    content = b"SAME" * 10
    (pending / "dup.pdf").write_bytes(content)
    review = Path(config.root_dir) / "非发票附件"
    review.mkdir(parents=True)
    (review / "dup.pdf").write_bytes(content)

    cand = ArchiveCandidate(
        record=_record(doc_type="travel_voucher", file_name="dup.pdf",
                       file_path=str(pending / "dup.pdf")),
        source_path=str(pending / "dup.pdf"),
        duplicate_status="新发票", duplicate_reason="",
        eligible=False, reasons=["not_e_invoice"],
    )
    result = service.archive_non_invoices([cand], config)
    assert len(result["merged"]) == 1
    assert not (pending / "dup.pdf").exists()
    assert len(list(review.glob("dup*.pdf"))) == 1


def test_archive_non_invoices_suffixes_diff_content(tmp_path: Path) -> None:
    service, config, pending = _service(tmp_path)
    (pending / "same.pdf").write_bytes(b"NEW")
    review = Path(config.root_dir) / "非发票附件"
    review.mkdir(parents=True)
    (review / "same.pdf").write_bytes(b"OLD")

    cand = ArchiveCandidate(
        record=_record(doc_type="hotel_folio", file_name="same.pdf",
                       file_path=str(pending / "same.pdf")),
        source_path=str(pending / "same.pdf"),
        duplicate_status="新发票", duplicate_reason="",
        eligible=False, reasons=["not_e_invoice"],
    )
    service.archive_non_invoices([cand], config)
    names = sorted(p.name for p in review.iterdir())
    assert names == ["same.pdf", "same_1.pdf"]
