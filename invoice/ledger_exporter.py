"""Monthly reimbursement ledger exporter (S3.3).

Two sheets per month file:

- ``报销流水`` — the printable reimbursement form mirroring the user's
  existing template (序号/日期/姓名/金额/类别/凭证类型/费用发生地/备注),
  plus 报销状态 so invoice archive ≠ reimbursed expense stays explicit.
  备注 is always blank by design.
- ``发票明细`` — full tax fields + local PDF path for traceability.

Export is a pure function of archived records, so regenerating a month
can never double-count.  Writes go through a temp file + os.replace.
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from models.invoice import InvoiceRecord

FLOW_HEADERS = ["序号", "日期", "姓名", "金额", "类别", "凭证类型", "费用发生地", "备注", "报销状态"]
DETAIL_HEADERS = [
    "序号", "发票号码", "开票日期", "价税合计", "税额", "购买方", "销售方",
    "发票内容", "文件名", "PDF路径", "归档状态", "审核状态", "重复状态", "来源",
]
_FLOW_WIDTHS = [6, 12, 10, 10, 12, 10, 14, 16, 10]
_DETAIL_WIDTHS = [6, 22, 12, 10, 10, 30, 30, 28, 36, 44, 12, 10, 10, 16]

_ARCHIVE_TEXT = {"auto_archived": "自动归档", "manually_archived": "人工确认归档", "deferred": "暂缓"}
_REVIEW_TEXT = {"pending_review": "待审核", "confirmed": "已确认", "": ""}
_REIMBURSE_TEXT = {"undecided": "未确定", "reimbursed": "报销", "excluded": "不报销"}
_PENDING = "待指定"
_PENDING_CATEGORY = "待确认"


class MonthlyLedgerExporter:
    """Write the two-sheet monthly ledger workbook."""

    def export(
        self,
        path: Path,
        records: list[InvoiceRecord],
        reimburser_default: str = "",
    ) -> Path:
        rows = sorted(
            (r for r in records if r.invoice_date),
            key=lambda r: (r.invoice_date, r.invoice_number or ""),
        )
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        self._write_flow_sheet(workbook.active, rows, reimburser_default)
        self._write_detail_sheet(workbook.create_sheet("发票明细"), rows)

        fd, tmp_name = tempfile.mkstemp(
            dir=str(path.parent), prefix=path.stem + ".", suffix=".tmp"
        )
        os.close(fd)
        try:
            workbook.save(tmp_name)
            os.replace(tmp_name, path)
        except BaseException:
            Path(tmp_name).unlink(missing_ok=True)
            raise
        return path

    def _write_flow_sheet(self, sheet, rows, reimburser_default: str) -> None:
        sheet.title = "报销流水"
        sheet.append(FLOW_HEADERS)
        total = 0.0
        for idx, record in enumerate(rows, start=1):
            total += record.amount or 0.0
            sheet.append(
                [
                    idx,
                    record.invoice_date.replace("-", "."),
                    record.reimburser or reimburser_default or _PENDING,
                    record.amount,
                    record.expense_category or _PENDING_CATEGORY,
                    "电子发票",
                    record.expense_location or "",
                    "",
                    _REIMBURSE_TEXT.get(record.reimburse_state, "未确定"),
                ]
            )
        sheet.append(["合计", "", "", round(total, 2), f"共 {len(rows)} 张", "", "", "", ""])
        self._finish_sheet(sheet, _FLOW_WIDTHS, last_row=len(rows) + 1)
        for row_idx in range(2, len(rows) + 3):
            cell = sheet.cell(row=row_idx, column=4)
            cell.number_format = "0.00"

    def _write_detail_sheet(self, sheet, rows) -> None:
        sheet.append(DETAIL_HEADERS)
        for idx, record in enumerate(rows, start=1):
            sheet.append(
                [
                    idx,
                    record.invoice_number,
                    record.invoice_date,
                    record.amount,
                    record.tax_amount,
                    record.buyer or "",
                    record.seller or "",
                    record.invoice_content or "",
                    record.file_name,
                    record.file_path,
                    _ARCHIVE_TEXT.get(record.archive_status, record.archive_status or ""),
                    _REVIEW_TEXT.get(record.review_status, record.review_status or ""),
                    "新发票",
                    record.source_type or "",
                ]
            )
        self._finish_sheet(sheet, _DETAIL_WIDTHS)
        for row_idx in range(2, len(rows) + 2):
            sheet.cell(row=row_idx, column=4).number_format = "0.00"
            sheet.cell(row=row_idx, column=5).number_format = "0.00"

    @staticmethod
    def _finish_sheet(sheet, widths: list[int], last_row: int | None = None) -> None:
        for idx, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
        sheet.freeze_panes = "A2"
        last = last_row or sheet.max_row
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{last}"
