"""Excel ledger export for invoice archives."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from models.invoice import InvoiceRecord

HEADERS = [
    "序号",
    "发票编号",
    "开票日期",
    "发票内容",
    "发票金额",
    "商户/销售方",
    "文件名",
    "文件路径",
    "查重结果",
    "识别状态",
    "日期来源",
    "备注",
    "税额",
    "购买方",
]


class InvoiceExcelExporter:
    """Write monthly invoice ledgers to XLSX files."""

    def export(self, report_path: Path, records: list[InvoiceRecord]) -> Path:
        """Create or overwrite a workbook for the provided invoice records."""
        report_path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "台账"

        for column_index, header in enumerate(HEADERS, start=1):
            sheet.cell(row=1, column=column_index, value=header)

        total_amount = 0.0
        for row_index, record in enumerate(records, start=2):
            amount = float(record.amount or 0.0)
            total_amount += amount

            sheet.cell(row=row_index, column=1, value=row_index - 1)
            sheet.cell(row=row_index, column=2, value=record.invoice_number)
            sheet.cell(row=row_index, column=3, value=record.invoice_date)
            sheet.cell(row=row_index, column=4, value=record.invoice_content)
            sheet.cell(row=row_index, column=5, value=amount)
            sheet.cell(row=row_index, column=6, value=record.seller)
            sheet.cell(row=row_index, column=7, value=record.file_name)
            sheet.cell(row=row_index, column=8, value=record.file_path)
            sheet.cell(row=row_index, column=9, value="新发票")
            sheet.cell(row=row_index, column=10, value=record.recognition_status)
            sheet.cell(row=row_index, column=11, value=record.date_source)
            sheet.cell(row=row_index, column=12, value="")
            sheet.cell(
                row=row_index,
                column=13,
                value=(
                    float(record.tax_amount)
                    if record.tax_amount is not None
                    else ""
                ),
            )
            sheet.cell(row=row_index, column=14, value=record.buyer)

        summary_row = len(records) + 3
        sheet.cell(row=summary_row, column=1, value="发票总数")
        sheet.cell(row=summary_row, column=2, value=len(records))
        sheet.cell(row=summary_row + 1, column=1, value="金额总计")
        sheet.cell(row=summary_row + 1, column=2, value=total_amount)

        workbook.save(report_path)
        return report_path
